import os
from flask import Flask, render_template, request, send_file, flash, redirect, url_for
from werkzeug.utils import secure_filename
import openpyxl
from openpyxl.styles import PatternFill

app = Flask(__name__)
app.secret_key = 'super-secret-key'

UPLOAD_FOLDER = os.path.join(os.path.dirname(__file__), 'uploads')
OUTPUT_FOLDER = os.path.join(os.path.dirname(__file__), 'outputs')

os.makedirs(UPLOAD_FOLDER, exist_ok=True)
os.makedirs(OUTPUT_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['OUTPUT_FOLDER'] = OUTPUT_FOLDER
app.config['ALLOWED_EXTENSIONS'] = {'xlsx', 'xls'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        if 'obsolete_file' not in request.files or 'target_file' not in request.files:
            flash('Please select both files before submitting.', 'error')
            return redirect(request.url)
        
        file1 = request.files['obsolete_file']
        file2 = request.files['target_file']
        
        if file1.filename == '' or file2.filename == '':
            flash('Both file inputs are required.', 'error')
            return redirect(request.url)

        # STRICT BACKEND CHECK: Verify both files have valid Excel extensions
        if not allowed_file(file1.filename) or not allowed_file(file2.filename):
            flash('Invalid file type! Please upload only Excel (.xlsx or .xls) files.', 'error')
            return redirect(url_for('index'))

        obs_filename = secure_filename(file1.filename)
        tgt_filename = secure_filename(file2.filename)
        
        obsolete_path = os.path.join(app.config['UPLOAD_FOLDER'], obs_filename)
        target_path = os.path.join(app.config['UPLOAD_FOLDER'], tgt_filename)
        
        file1.save(obsolete_path)
        file2.save(target_path)

        color_hex = request.form.get('highlight_color', '#ADD8E6').lstrip('#')

        try:
            wb_obs = openpyxl.load_workbook(obsolete_path, data_only=True)
            sheet_obs = wb_obs.active
            
            obsolete_parts = set()
            for row in sheet_obs.iter_rows(min_col=1, max_col=1, values_only=True):
                val = row[0]
                if val is not None:
                    obsolete_parts.add(str(val).strip())

            wb_target = openpyxl.load_workbook(target_path)
            sheet_target = wb_target.active

            custom_fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")

            match_count = 0
            for cell in sheet_target['A']:
                if cell.value is not None:
                    cell_val = str(cell.value).strip()
                    if cell_val in obsolete_parts:
                        cell.fill = custom_fill
                        match_count += 1

            output_filename = f"highlighted_{tgt_filename}"
            output_filepath = os.path.join(app.config['OUTPUT_FOLDER'], output_filename)
            wb_target.save(output_filepath)

        except Exception as e:
            flash(f'An error occurred while processing the Excel files: {e}', 'error')
            return redirect(url_for('index'))

        finally:
            if os.path.exists(obsolete_path):
                os.remove(obsolete_path)
            if os.path.exists(target_path):
                os.remove(target_path)

        return render_template('index.html', download_file=output_filename, matches=match_count)

    return render_template('index.html', download_file=None)

@app.route('/download/<filename>')
def download_file(filename):
    file_path = os.path.join(app.config['OUTPUT_FOLDER'], filename)
    
    if not os.path.exists(file_path):
        return redirect(url_for('index'))

    response = send_file(file_path, as_attachment=True)
    
    @response.call_on_close
    def cleanup():
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
        except Exception as e:
            app.logger.error(f"Error removing file: {e}")

    return response

if __name__ == '__main__':
    app.run(debug=False)